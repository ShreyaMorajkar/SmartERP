import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Colors
    BG_COLOR = RGBColor(15, 23, 42)      # Slate 900
    CARD_BG = RGBColor(30, 41, 59)       # Slate 800
    CARD_BORDER = RGBColor(51, 65, 85)   # Slate 700
    TEXT_MAIN = RGBColor(248, 250, 252)  # Slate 50
    TEXT_MUTED = RGBColor(148, 163, 184) # Slate 400
    ACCENT_GREEN = RGBColor(16, 185, 129)# Emerald 500
    ACCENT_CYAN = RGBColor(6, 182, 212)  # Cyan 500
    ACCENT_RED = RGBColor(239, 68, 68)   # Red 500

    def apply_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.color.rgb = BG_COLOR

    def add_header(slide, title_text, category="SMARTERP PRESENTATION"):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.7), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        r_cat = p_cat.add_run()
        r_cat.text = category.upper()
        r_cat.font.name = 'Arial'
        r_cat.font.size = Pt(10)
        r_cat.font.bold = True
        r_cat.font.color.rgb = ACCENT_GREEN

        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(0.65))
        tf_t = t_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
        p_t = tf_t.paragraphs[0]
        r_t = p_t.add_run()
        r_t.text = title_text
        r_t.font.name = 'Arial'
        r_t.font.size = Pt(22)
        r_t.font.bold = True
        r_t.font.color.rgb = TEXT_MAIN

    def add_card(slide, left, top, width, height, title, items, badge_text=None, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = border_color
        card.line.width = Pt(1)

        tb = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.25), width - Inches(0.5), height - Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_title = tf.paragraphs[0]
        r_title = p_title.add_run()
        r_title.text = title
        r_title.font.name = 'Arial'
        r_title.font.size = Pt(14)
        r_title.font.bold = True
        r_title.font.color.rgb = ACCENT_CYAN

        if badge_text:
            r_badge = p_title.add_run()
            r_badge.text = f"  [{badge_text}]"
            r_badge.font.name = 'Arial'
            r_badge.font.size = Pt(10)
            r_badge.font.bold = True
            r_badge.font.color.rgb = ACCENT_GREEN

        p_title.space_after = Pt(10)

        for item in items:
            p_item = tf.add_paragraph()
            p_item.space_after = Pt(6)
            if isinstance(item, tuple):
                bold_txt, regular_txt = item
                r_b = p_item.add_run()
                r_b.text = "• " + bold_txt + ": "
                r_b.font.name = 'Arial'
                r_b.font.size = Pt(11)
                r_b.font.bold = True
                r_b.font.color.rgb = TEXT_MAIN
                
                r_r = p_item.add_run()
                r_r.text = regular_txt
                r_r.font.name = 'Arial'
                r_r.font.size = Pt(11)
                r_r.font.color.rgb = TEXT_MUTED
            else:
                r_text = p_item.add_run()
                r_text.text = "• " + item
                r_text.font.name = 'Arial'
                r_text.font.size = Pt(11)
                r_text.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 1: Title Slide (Simple & Clear)
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    apply_background(slide1)

    tag_box = slide1.shapes.add_textbox(Inches(1.2), Inches(1.5), Inches(10.9), Inches(0.4))
    tf_tag = tag_box.text_frame
    p_tag = tf_tag.paragraphs[0]
    r_tag = p_tag.add_run()
    r_tag.text = "ACADEMIC PROJECT PRESENTATION (3-4 MINS)"
    r_tag.font.name = 'Arial'
    r_tag.font.size = Pt(12)
    r_tag.font.bold = True
    r_tag.font.color.rgb = ACCENT_GREEN

    t_box1 = slide1.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(10.9), Inches(1.2))
    tf_t1 = t_box1.text_frame
    p_t1 = tf_t1.paragraphs[0]
    r_t1 = p_t1.add_run()
    r_t1.text = "SmartERP"
    r_t1.font.name = 'Arial'
    r_t1.font.size = Pt(44)
    r_t1.font.bold = True
    r_t1.font.color.rgb = TEXT_MAIN

    sub_box1 = slide1.shapes.add_textbox(Inches(1.2), Inches(3.2), Inches(10.9), Inches(0.8))
    tf_sub1 = sub_box1.text_frame
    p_sub1 = tf_sub1.paragraphs[0]
    r_sub1 = p_sub1.add_run()
    r_sub1.text = "A Fast, Keyboard-Friendly Web App for Billing, Stock & Accounting"
    r_sub1.font.name = 'Arial'
    r_sub1.font.size = Pt(18)
    r_sub1.font.color.rgb = ACCENT_CYAN

    info_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(4.3), Inches(10.9), Inches(2.0))
    info_card.fill.solid()
    info_card.fill.fore_color.rgb = CARD_BG
    info_card.line.color.rgb = CARD_BORDER

    tf_info = info_card.text_frame
    tf_info.word_wrap = True
    
    p_pres = tf_info.paragraphs[0]
    r_p1 = p_pres.add_run()
    r_p1.text = "Presented by: "
    r_p1.font.bold = True
    r_p1.font.size = Pt(14)
    r_p1.font.color.rgb = TEXT_MAIN
    r_p2 = p_pres.add_run()
    r_p2.text = "Shreya Morajkar  |  "
    r_p2.font.size = Pt(14)
    r_p2.font.color.rgb = ACCENT_GREEN

    r_p3 = p_pres.add_run()
    r_p3.text = "Project: "
    r_p3.font.bold = True
    r_p3.font.size = Pt(14)
    r_p3.font.color.rgb = TEXT_MAIN
    r_p4 = p_pres.add_run()
    r_p4.text = "Full-Stack Web ERP Application"
    r_p4.font.size = Pt(14)
    r_p4.font.color.rgb = TEXT_MUTED

    p_tech = tf_info.add_paragraph()
    p_tech.space_before = Pt(10)
    r_tc1 = p_tech.add_run()
    r_tc1.text = "Technologies Used: "
    r_tc1.font.bold = True
    r_tc1.font.size = Pt(13)
    r_tc1.font.color.rgb = TEXT_MAIN
    r_tc2 = p_tech.add_run()
    r_tc2.text = "Next.js (React), Express.js (Node.js backend), Prisma ORM & SQLite Database"
    r_tc2.font.size = Pt(13)
    r_tc2.font.color.rgb = TEXT_MUTED

    slide1.notes_slide.notes_text_frame.text = (
        "WHAT TO SAY (Slide 1 - 25 seconds):\n"
        "Good morning/afternoon! Today I am presenting my project called SmartERP.\n"
        "It is a full-stack web application designed for billing, inventory, and accounting.\n"
        "The special feature of this project is that it is keyboard-friendly — inspired by Tally — so accountants can enter bills without constantly needing to use the mouse.\n"
        "I built the frontend in Next.js and React, and the backend in Node.js, Express, and Prisma."
    )

    # ==========================================
    # SLIDE 2: Why We Built This (Problem & Solution)
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    apply_background(slide2)
    add_header(slide2, "Why Did We Build SmartERP?", "BACKGROUND & MOTIVATION")

    add_card(slide2, Inches(0.8), Inches(1.6), Inches(3.7), Inches(5.2), "The Problem Today", [
        ("Too Many Mouse Clicks", "Most modern web software requires clicking on multiple dropdowns, which is slow for typing lots of bills."),
        ("Desktop Tally is Offline", "Old Tally software is fast, but it only runs on one computer and cannot be easily used over the web."),
        ("Math Mistakes", "If billing, stock, and accounts are in different places, numbers get mixed up.")
    ], "THE ISSUE", ACCENT_RED)

    add_card(slide2, Inches(4.8), Inches(1.6), Inches(3.7), Inches(5.2), "The Idea / Inspiration", [
        ("Tally's Keyboard Speed", "Keep the fast keyboard shortcuts that accountants love (like F8 for Sales, F9 for Purchase, F4 for Calculator)."),
        ("Anywhere Web Access", "Run in any browser so business owners can open it from anywhere."),
        ("Multi-Company Support", "One user can easily switch between up to 5 different businesses.")
    ], "THE INSPIRATION", ACCENT_CYAN)

    add_card(slide2, Inches(8.8), Inches(1.6), Inches(3.7), Inches(5.2), "Our SmartERP Solution", [
        ("100% Keyboard Workflows", "Create masters, record sales, and view reports just using keyboard keys."),
        ("Automatic Calculations", "System automatically calculates GST tax (CGST, SGST, IGST) and updates stock count."),
        ("Accurate Books", "Every entry updates Debits and Credits together so books always balance.")
    ], "THE SOLUTION", ACCENT_GREEN)

    slide2.notes_slide.notes_text_frame.text = (
        "WHAT TO SAY (Slide 2 - 35 seconds):\n"
        "Why did we build this?\n"
        "Accountants enter hundreds of bills every day. Normal websites force them to click around with a mouse, which slows them down.\n"
        "On the other hand, traditional desktop accounting software like Tally is fast on keyboard, but it is locked to a single PC.\n"
        "SmartERP gives you the best of both worlds: it runs smoothly on the web in any browser, but lets you type bills quickly using standard keyboard shortcuts."
    )

    # ==========================================
    # SLIDE 3: How the System Works (Simple Architecture)
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    apply_background(slide3)
    add_header(slide3, "How the System is Built (Architecture)", "TECH STACK")

    add_card(slide3, Inches(0.8), Inches(1.6), Inches(3.7), Inches(5.2), "1. Frontend (User Interface)", [
        ("Next.js & React", "Renders the web pages quickly and handles all user interactions."),
        ("Tailwind CSS", "Provides a clean, dark-mode design with highlighted shortcut keys."),
        ("Keyboard Engine", "Listens for keys like F1 to F9, Alt, and Ctrl to switch screens instantly."),
        ("Built-in Calculator", "Press F4 anywhere to open a popup calculator without losing work.")
    ], "FRONTEND")

    add_card(slide3, Inches(4.8), Inches(1.6), Inches(3.7), Inches(5.2), "2. Backend (API Server)", [
        ("Node.js & Express", "Handles all incoming requests, login security, and math logic."),
        ("JWT Security", "Users log in securely and passwords are encrypted using bcrypt."),
        ("Tax Engine", "Checks company state vs customer state to decide if CGST+SGST or IGST applies."),
        ("Export Engine", "Generates PDF Tax Invoices and Excel spreadsheets for reports.")
    ], "BACKEND API")

    add_card(slide3, Inches(8.8), Inches(1.6), Inches(3.7), Inches(5.2), "3. Database & ORM", [
        ("Prisma ORM", "Connects our Node.js server to the database with clean code and zero SQL errors."),
        ("Safe Transactions", "When you save a bill, it updates customer balance and stock quantity at the exact same time."),
        ("SQLite Database", "Stores companies, accounts, stock items, and bills reliably in one place.")
    ], "DATABASE")

    slide3.notes_slide.notes_text_frame.text = (
        "WHAT TO SAY (Slide 3 - 40 seconds):\n"
        "Here is the simple technical architecture:\n"
        "1. On the frontend, we use Next.js and React. We wrote a custom keyboard engine that listens to keys like F8, F9, and F4.\n"
        "2. The frontend talks to a Node.js and Express backend API. The backend verifies user logins, handles business rules, and computes GST taxes.\n"
        "3. For the database, we use Prisma ORM and SQLite. When a user creates a bill, Prisma runs an 'atomic transaction' — meaning it updates the customer balance and lowers the stock quantity at the exact same moment so data never gets corrupted."
    )

    # ==========================================
    # SLIDE 4: Core Features & Easy Workflow
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    apply_background(slide4)
    add_header(slide4, "Main Modules & Daily Workflow", "HOW A USER WORKS")

    add_card(slide4, Inches(0.8), Inches(1.6), Inches(5.7), Inches(2.45), "1. Automatic Company Setup", [
        ("Instant Seeding", "Creating a company automatically creates 4 base accounting groups (Assets, Liabilities, Income, Expenses) and Cash/Bank accounts."),
        ("Multiple Companies", "Switch between companies easily using the F1 key.")
    ], "STEP 1")

    add_card(slide4, Inches(6.8), Inches(1.6), Inches(5.7), Inches(2.45), "2. Creating Masters (Ledgers & Stock)", [
        ("Accounts", "Create Suppliers, Customers, Sales A/c, and Purchase A/c in the Masters panel (Press M)."),
        ("Products", "Add items with unit (PCS), purchase price, selling price, and GST % (e.g. 18%).")
    ], "STEP 2")

    add_card(slide4, Inches(0.8), Inches(4.3), Inches(5.7), Inches(2.5), "3. Recording Bills (Vouchers)", [
        ("Purchase (F9)", "Buy stock from suppliers $\\rightarrow$ stock increases, supplier credit balance increases."),
        ("Sales (F8)", "Sell stock to customers $\\rightarrow$ stock decreases, customer debit balance increases, GST calculated automatically.")
    ], "STEP 3")

    add_card(slide4, Inches(6.8), Inches(4.3), Inches(5.7), Inches(2.5), "4. Invoices & Downloads", [
        ("PDF Invoices", "Click to download a neat, professional PDF invoice with tax breakdown."),
        ("Excel Reports", "Click 'Export to Excel' to download any accounting statement.")
    ], "STEP 4")

    slide4.notes_slide.notes_text_frame.text = (
        "WHAT TO SAY (Slide 4 - 45 seconds):\n"
        "How does a user actually work on SmartERP?\n"
        "Step 1: They create their company. The system automatically sets up default cash and bank accounts.\n"
        "Step 2: They add their customers, suppliers, and products with GST rates.\n"
        "Step 3: They record transactions. Pressing F9 opens the Purchase Voucher to buy stock. Pressing F8 opens the Sales Voucher to sell items. The system automatically calculates 18% GST and adjusts the inventory count.\n"
        "Step 4: They can print clean PDF tax invoices or export Excel sheets."
    )

    # ==========================================
    # SLIDE 5: Real-Time Accounting Reports
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    apply_background(slide5)
    add_header(slide5, "Real-Time Reports & Verification", "ACCOUNTING ACCURACY")

    add_card(slide5, Inches(0.8), Inches(1.6), Inches(3.7), Inches(5.2), "Instant Accounting Statements", [
        ("Trial Balance (Alt+T)", "Shows all accounts. Total Debits and Total Credits match 100% in real time."),
        ("Profit & Loss (Alt+P)", "Shows Total Revenue minus Total Expenses to calculate Net Profit or Loss."),
        ("Balance Sheet (Alt+B)", "Shows Total Assets equal Total Liabilities."),
        ("Stock Summary (Alt+R)", "Shows remaining quantities and total valuation of warehouse inventory.")
    ], "STATEMENTS")

    add_card(slide5, Inches(4.8), Inches(1.6), Inches(3.7), Inches(5.2), "GST Tax Summary (Alt+X)", [
        ("CGST & SGST", "Automatically tracked when buying/selling within the same state."),
        ("IGST", "Automatically tracked for sales to other states."),
        ("Tax Summary", "Shows total tax collected vs total tax paid so businesses know how much GST to pay.")
    ], "TAX REGISTERS")

    add_card(slide5, Inches(8.8), Inches(1.6), Inches(3.7), Inches(5.2), "Automated Testing Proof", [
        ("Built-in Test Harness", "We created an automated test script (test_features.js)."),
        ("Tests All Modules", "It tests Signup, Company Creation, Purchase, Sale, and Math calculations automatically."),
        ("100% Passed", "Guarantees that all double-entry arithmetic is mathematically correct.")
    ], "VERIFIED", ACCENT_GREEN)

    slide5.notes_slide.notes_text_frame.text = (
        "WHAT TO SAY (Slide 5 - 40 seconds):\n"
        "One of the biggest strengths of SmartERP is that all accounting reports update live in real time:\n"
        "1. The Trial Balance instantly verifies that all Debits match Credits.\n"
        "2. The Profit & Loss shows revenue, expenses, and net profit.\n"
        "3. The Balance Sheet verifies assets and liabilities.\n"
        "4. And the GST Register calculates exact tax liabilities.\n"
        "To make sure there are zero mathematical bugs, we also wrote an automated test script (test_features.js) that tests every single formula automatically."
    )

    # ==========================================
    # SLIDE 6: Keyboard-First Shortcuts & Productivity
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    apply_background(slide6)
    add_header(slide6, "Keyboard Shortcuts for High Speed", "USER EXPERIENCE")

    add_card(slide6, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.2), "Simple Keyboard Hotkeys", [
        ("F1", "Switch Company"),
        ("F4", "Open / Close Floating Calculator"),
        ("F8", "Open Sales Voucher (for selling items)"),
        ("F9", "Open Purchase Voucher (for buying items)"),
        ("Ctrl + K", "Search anything across the app (Command Palette)"),
        ("ESC", "Go back to the previous screen or Gateway"),
        ("Red Underlined Letters", "Type 'M' for Masters, 'V' for Vouchers, 'T' for Trial Balance directly!")
    ], "KEYBOARD MAP")

    add_card(slide6, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), "Why Accountants Love This", [
        ("Popup Calculator (F4)", "Can be opened anytime while typing a bill to calculate discounts or taxes, and closed with one key without losing entered data."),
        ("Quick Search (Ctrl+K)", "Type 'Balance' or 'Sale' to jump straight to the page you want in 1 second."),
        ("No Mouse Needed", "Users can do 100% of their daily work keeping their hands on the keyboard, making billing 3 times faster.")
    ], "BENEFITS", ACCENT_CYAN)

    slide6.notes_slide.notes_text_frame.text = (
        "WHAT TO SAY (Slide 6 - 30 seconds):\n"
        "To make billing fast, we added keyboard shortcuts inspired by Tally:\n"
        "- Pressing F8 opens Sales, F9 opens Purchase.\n"
        "- Pressing F4 opens a popup calculator right on top of your screen to calculate taxes, and pressing F4 again closes it without resetting your form.\n"
        "- Pressing Ctrl+K opens a search bar to jump anywhere.\n"
        "This allows users to operate the entire application using only the keyboard."
    )

    # ==========================================
    # SLIDE 7: Conclusion & Summary
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    apply_background(slide7)
    add_header(slide7, "Conclusion & Project Summary", "SUMMARY")

    add_card(slide7, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.2), "What We Achieved", [
        ("Working Full-Stack ERP", "Successfully built a complete billing, inventory, and double-entry accounting software."),
        ("Speed & Ease of Use", "Combined modern web cloud tech with Tally's fast keyboard shortcuts."),
        ("Zero Calculation Errors", "Accurate double-entry math, automatic GST calculation, and real-time reports."),
        ("Real-World Ready", "Supports multi-company, PDF invoices, and Excel exports.")
    ], "ACHIEVEMENTS", ACCENT_GREEN)

    add_card(slide7, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), "Project Deliverables", [
        ("GitHub Repository", "Fully uploaded on GitHub (ShreyaMorajkar/SmartERP) with complete code."),
        ("Complete Documentation", "README with 14 real-time application screenshots showing all live views."),
        ("Automated Test Suite", "Integration tests verifying 100% correctness of accounting calculations."),
        ("Live Demo Available", "Ready to demonstrate live on localhost anytime!")
    ], "DELIVERABLES", ACCENT_CYAN)

    slide7.notes_slide.notes_text_frame.text = (
        "WHAT TO SAY (Slide 7 - 20 seconds):\n"
        "To conclude, SmartERP is a complete, working cloud accounting software that is fast, accurate, and easy to use.\n"
        "The entire project is pushed to GitHub with detailed documentation, real screenshots, and automated test cases.\n"
        "Thank you so much! I am happy to show a live demo or answer any questions."
    )

    # Save to Downloads
    downloads_path = os.path.expanduser('~/Downloads')
    file_path = os.path.join(downloads_path, 'SmartERP_Presentation.pptx')
    prs.save(file_path)
    print(f"Simplified authentic presentation successfully created and saved to: {file_path}")

if __name__ == '__main__':
    create_presentation()
