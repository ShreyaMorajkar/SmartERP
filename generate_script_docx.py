import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_document():
    doc = docx.Document()
    
    # Set Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Style Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("SmartERP: 7-Minute Demonstration Video Script")
    title_run.font.name = 'Arial'
    title_run.font.size = Pt(18)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(16, 185, 129) # Emerald Green

    # Subtitle
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("Keyboard-Driven, Tally-Inspired Accounting & Billing Platform")
    sub_run.font.name = 'Arial'
    sub_run.font.size = Pt(11)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(100, 116, 139) # Slate Grey

    doc.add_paragraph().paragraph_format.space_after = Pt(20)

    # Table of Contents / Outline
    h_outline = doc.add_paragraph()
    h_outline_run = h_outline.add_run("Video Script Timeline & Structure")
    h_outline_run.font.name = 'Arial'
    h_outline_run.font.size = Pt(14)
    h_outline_run.font.bold = True
    h_outline_run.font.color.rgb = RGBColor(30, 41, 59)

    outline_table = doc.add_table(rows=1, cols=3)
    outline_table.style = 'Light Shading Accent 1'
    hdr_cells = outline_table.rows[0].cells
    hdr_cells[0].text = 'Time'
    hdr_cells[1].text = 'Segment'
    hdr_cells[2].text = 'Goal'

    segments = [
        ("0:00 - 0:45", "Part 1: Intro & Technology Stack", "Introduce SmartERP, its tech stack (Next.js, Express, Prisma, SQLite), and its keyboard-first design."),
        ("0:45 - 1:45", "Part 2: Login & Company Seeding", "Show administrator signup/login and company creation. Detail the automatic database group/ledger seeding."),
        ("1:45 - 3:00", "Part 3: Masters Setup", "Set up ledgers for customers/suppliers and create a stock item product with 18% GST."),
        ("3:00 - 4:15", "Part 4: Purchase Transaction", "Record a credit purchase voucher using the F9 hotkey. Verify stock level increment and account balances."),
        ("4:15 - 5:30", "Part 5: Sales Transaction", "Record a credit sales voucher using the F8 hotkey. Show automatic GST CGST/SGST tax split calculation."),
        ("5:30 - 6:30", "Part 6: Reports & Document Exports", "Verify balanced books in Trial Balance and inspect Profit & Loss. Download a PDF invoice and Excel export."),
        ("6:30 - 7:00", "Part 7: Keyboard Shortcuts & Outro", "Highlight keyboard efficiency features like the F4 Calculator, Ctrl+K Command Palette, and menu letter hotkeys.")
    ]

    for time, seg, goal in segments:
        row_cells = outline_table.add_row().cells
        row_cells[0].text = time
        row_cells[1].text = seg
        row_cells[2].text = goal

    doc.add_paragraph().paragraph_format.space_after = Pt(20)

    # Detailed Script Sections
    script_data = [
        {
            "title": "0:00 - 0:45 | Part 1: Intro & Technology Stack (45 seconds)",
            "visual": "A clean desktop viewport displaying the dashboard login page at http://localhost:3030. The interface shows a dark-themed glassmorphic card glowing against a slate-blue background. The presenter is visible in a small circular webcam crop in the bottom corner.",
            "action": "Select the email login field and type the admin username.",
            "narration": "Hello everyone! Today, I’m excited to show you SmartERP—a cloud-based, keyboard-first Billing, Inventory, and Accounting Management System designed specifically for accountants. Unlike traditional mouse-heavy ERPs, SmartERP is keyboard-driven, inspired by Tally, and built using Next.js, Express, and SQLite. It delivers lightning-fast data entry for high-volume transactions while preserving modern cloud scalability. Let’s log in and initialize our first workspace."
        },
        {
            "title": "0:45 - 1:45 | Part 2: Login & Company Seeding (60 seconds)",
            "visual": "The viewport transitions to the User Sign Up form, submits, and loads the Company Selection workspace panel. Presenter triggers the 'Initialize New Company' dialog.",
            "action": "1. Click 'Sign Up'. Input Name: 'Demo Admin', Email: 'admin@smarterp.com', Password: 'admin123'. Click 'Register Admin'.\n2. Click 'Initialize New Company'. Type Name: 'Acme Trading Solutions', State: 'Goa', GSTIN: '30AAAAA1111A1Z1'. Click 'Submit'.\n3. Click on the company card to enter the dashboard.",
            "narration": "First, we register our administrator account. Logging in brings us to the Company Selection Screen where users can manage up to 5 distinct businesses. Let's create a new company—Acme Trading Solutions. As soon as this company is initialized, the Express backend automatically seeds 4 core accounting groups—Assets, Liabilities, Income, and Expenses—and 2 default system accounts—Cash and Bank. This means we are ready to write books immediately."
        },
        {
            "title": "1:45 - 3:00 | Part 3: Masters Setup (75 seconds)",
            "visual": "The Gateway of SmartERP dashboard. The presenter navigates to the Masters creation panel.",
            "action": "1. Press 'M' on the keyboard (observe the red underlined 'M' on 'Create/Alter Masters').\n2. Create a ledger named 'Sun Distributors' under Liabilities. GSTIN: '30BBBBB2222B2Z2'. Submit.\n3. Create a ledger named 'Ocean Retailers' under Assets. GSTIN: '30CCCCC3333C3Z3'. Submit.\n4. Create ledgers 'Purchase A/c' (Expenses) and 'Sales A/c' (Income).\n5. Go to Stock Item tab, Name: 'Smart LED TV', SKU: 'TV-001', Purchase Rate: '10000', Selling Rate: '15000', GST: '18%', Opening Qty: '100'. Submit.",
            "narration": "Now we are at the Gateway of SmartERP. Notice the red underlined letters on each option. Pressing 'M' takes us instantly to the Masters Creation Panel. Let's set up our custom ledgers. We will create a supplier ledger, Sun Distributors, under Liabilities, and a customer, Ocean Retailers, under Assets. We also create standard Purchase and Sales accounts. Finally, we register our product: the 'Smart LED TV' priced at 10,000 INR purchase and 15,000 INR sale, with an 18% GST tax rate. The masters are now complete."
        },
        {
            "title": "3:00 - 4:15 | Part 4: Purchase Transaction (75 seconds)",
            "visual": "Voucher Entry spreadsheet grid view.",
            "action": "1. Press 'ESC' to return to Gateway, then 'V' to open Vouchers.\n2. Press 'F9' to open the Purchase Voucher entry screen.\n3. Select Party Ledger: 'Sun Distributors'.\n4. Under Inventory, select 'Smart LED TV', set quantity: '20', rate: '10000'. Notice the double entry balanced at the bottom. Click 'Submit'.",
            "narration": "To record a purchase, we navigate back to the Gateway and press 'V' to open Vouchers. Hitting 'F9' immediately switches our layout to the Purchase Voucher. We select Sun Distributors as our supplier. Under the inventory section, we select the Smart LED TV and enter a quantity of 20. The system calculates a base total of 200,000 INR. Because both the company and the supplier are registered in Goa, the system automatically applies local GST, splitting it into 18,000 INR CGST and 18,000 INR SGST. At the bottom, the double-entry ledger allocations are balanced automatically. Let's submit this voucher."
        },
        {
            "title": "4:15 - 5:30 | Part 5: Sales Transaction (75 seconds)",
            "visual": "Sales Voucher spreadsheet grid view.",
            "action": "1. Press 'F8' to switch to the Sales Voucher entry screen.\n2. Select Party Ledger: 'Ocean Retailers'.\n3. Under Inventory, select 'Smart LED TV', set quantity: '10'. Click 'Submit'.",
            "narration": "Now we will record a sale. Pressing 'F8' instantly loads the Sales Voucher layout. We select Ocean Retailers as our customer and add the Smart LED TV. Notice that the selling rate of 15,000 INR is auto-populated. We input a quantity of 10. The base value is 150,000 INR, and with the 18% GST tax, the final invoice total is 177,000 INR. The balanced ledger allocations show a Debit to the customer's account and a Credit to our Sales control account. Let's submit the sale. The inventory ledger automatically decrements the stock by 10 units."
        },
        {
            "title": "5:30 - 6:30 | Part 6: Reports & Document Exports (60 seconds)",
            "visual": "The Reports dashboard. Toggle tabs for Trial Balance, P&L, Balance Sheet, and GSTR register.",
            "action": "1. Press 'ESC' to return to Gateway, then 'T' to open Trial Balance.\n2. Press 'Alt + P' to show the Profit & Loss report.\n3. Press 'Alt + X' to open the GST Register.\n4. Click 'PDF Invoice' on the sales row to trigger the invoice pdf download.\n5. Click 'Export to Excel' to download the spreadsheet.",
            "narration": "With transactions posted, all ledger books update in real-time. By navigating to Reports, we see our Trial Balance is perfectly synced at 413,000 INR. Our Profit & Loss sheet details the net margins, and the Balance Sheet balances exactly. Under the GST Tax Register, we can see GSTR records. We can click 'PDF Invoice' to download a clean, tax-compliant invoice generated by PDFKit. We can also click 'Export to Excel' to instantly get an Excel file containing the full Trial Balance ledger."
        },
        {
            "title": "6:30 - 7:00 | Part 7: Keyboard Shortcuts & Outro (30 seconds)",
            "visual": "Presenter demonstrates F4 calculator overlay and command palette.",
            "action": "1. Press 'F4' to bring up the floating calculator. Type '15000 * 1.18' and hit enter. Press 'F4' to close.\n2. Press 'Ctrl + K' to trigger the Command Palette. Type 'Select Company', press 'Enter' to return to the company screen.",
            "narration": "Finally, let's look at the keyboard navigation. Pressing 'F4' opens a floating calculator overlay anywhere in the app to check margins. Pressing 'Ctrl + K' opens the Command Palette, letting us jump to any report or page instantly by typing search keys. SmartERP delivers a modern, keyboard-only cloud application that speeds up daily operations. Thank you for watching!"
        }
    ]

    for item in script_data:
        doc.add_heading(item["title"], level=2)
        
        # Format Visual
        p_vis = doc.add_paragraph()
        run_vis_hdr = p_vis.add_run("Visual Block: ")
        run_vis_hdr.bold = True
        run_vis_hdr.font.color.rgb = RGBColor(14, 165, 233) # Sky Blue
        p_vis.add_run(item["visual"])
        
        # Format Action
        p_act = doc.add_paragraph()
        run_act_hdr = p_act.add_run("Presenter Action: ")
        run_act_hdr.bold = True
        run_act_hdr.font.color.rgb = RGBColor(239, 68, 68) # Tally Red
        p_act.add_run(item["action"])

        # Format Narration
        p_narr = doc.add_paragraph()
        run_narr_hdr = p_narr.add_run("Narration: ")
        run_narr_hdr.bold = True
        run_narr_hdr.font.color.rgb = RGBColor(16, 185, 129) # Emerald Green
        
        run_narr_text = p_narr.add_run(f'"{item["narration"]}"')
        run_narr_text.italic = True
        
        # Spacing
        doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Save to Downloads
    downloads_path = os.path.expanduser('~/Downloads')
    file_path = os.path.join(downloads_path, 'SmartERP_Demo_Script.docx')
    doc.save(file_path)
    print(f"File successfully created and saved to: {file_path}")

if __name__ == '__main__':
    create_document()
